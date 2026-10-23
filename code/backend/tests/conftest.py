"""功能测试公共 fixtures 与工具函数。

数据环境约定（2026-10-22 与用户确认）：
- 复用开发库 mig_common_db 作为公共库；
- 每个测试会话创建独立测试空间（space_code=functest_*），空间库随之新建，
  会话结束时硬删除测试空间/用户及其空间库，保证可重复执行；
- 所有业务错误经全局异常处理器返回 HTTP 200 + 业务 code，断言请用 assert_biz_error。
"""

import csv
import io
import time
import uuid

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from main import app

API = "/api/v1"
ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "Admin@1234"
MEMBER_PASSWORD = "Member@1234"
VIEW_PASSWORD = "View@1234"

# 业务错误码（app/core/exceptions.py）
CODE_BIZ = 40000
CODE_CAPTCHA = 40001
CODE_PARAM = 42200
CODE_NOT_FOUND = 40400
CODE_FORBIDDEN = 40300
CODE_UNAUTHORIZED = 40100


# ── 基础断言工具 ──
def assert_ok(resp, expect_status: int = 200):
    """断言接口成功（HTTP 状态 + code=0），返回响应体。"""
    assert resp.status_code == expect_status, f"HTTP {resp.status_code}: {resp.text}"
    body = resp.json()
    assert body["code"] == 0, f"业务错误: {body}"
    return body


def assert_biz_error(resp, code: int, keyword: str | None = None):
    """断言业务错误：HTTP 200 + 指定业务 code（可附带消息关键字）。"""
    assert resp.status_code == 200, f"HTTP {resp.status_code}: {resp.text}"
    body = resp.json()
    assert body["code"] == code, f"期望 code={code}，实际: {body}"
    if keyword:
        assert keyword in body["message"], f"消息缺少「{keyword}」: {body['message']}"
    return body


# ── 文件构造工具（b03/b04 导入） ──
def xlsx_upload(headers: list, rows: list, filename: str = "import.xlsx") -> dict:
    """用 openpyxl 构造 xlsx 上传文件。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": (filename, buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def csv_upload(headers: list, rows: list, filename: str = "import.csv") -> dict:
    """构造 csv 上传文件（utf-8-sig，兼容 Excel）。"""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return {"file": (filename, buf.getvalue().encode("utf-8-sig"), "text/csv")}


def read_xlsx(content: bytes) -> list:
    """解析 xlsx 字节流为二维列表（含表头行）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def uid(prefix: str) -> str:
    """生成唯一名称（避免开发库残留数据干扰）。"""
    return f"{prefix}_{uuid.uuid4().hex[:8]}"


# ── 核心 fixtures ──
@pytest.fixture(scope="session")
def client():
    """会话级 TestClient：触发一次启动引导（建表/种子管理员）。"""
    with TestClient(app) as c:
        yield c


def _login(client, username: str, password: str, **extra) -> dict:
    return client.post(
        f"{API}/common/auth/login",
        json={"username": username, "password": password, **extra},
    ).json()


@pytest.fixture(scope="session")
def admin_headers(client):
    """系统管理员登录态（先复位失败计数，避免历史失败锁定）。"""
    from app.database import SessionLocal
    from app.models.user import CommonUser

    db = SessionLocal()
    try:
        admin = db.query(CommonUser).filter(CommonUser.username == ADMIN_USERNAME).first()
        if admin:
            admin.login_fail_count = 0
            admin.locked_until = None
            db.commit()
    finally:
        db.close()
    body = _login(client, ADMIN_USERNAME, ADMIN_PASSWORD)
    assert body["code"] == 0, f"管理员登录失败: {body}"
    return {"Authorization": f"Bearer {body['data']['token']}"}


def _create_user(client, admin_headers, username: str, password: str = MEMBER_PASSWORD) -> int:
    resp = client.post(
        f"{API}/common/users",
        headers=admin_headers,
        json={"username": username, "name": "功能测试用户", "org_type": "内部", "password": password},
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["code"] == 0, body
    return body["data"]["user"]["id"]


def _hard_delete_user(user_id: int) -> None:
    from app.database import SessionLocal
    from app.models.role import CommonUserRole
    from app.models.user import CommonUser

    db = SessionLocal()
    try:
        db.query(CommonUserRole).filter(CommonUserRole.user_id == user_id).delete()
        db.query(CommonUser).filter(CommonUser.id == user_id).delete()
        db.commit()
    finally:
        db.close()


@pytest.fixture(scope="session")
def test_space(client, admin_headers):
    """会话级测试空间：创建即自动建空间库与预置分类；结束即彻底清理。"""
    suffix = f"{int(time.time()) % 1000000}_{uuid.uuid4().hex[:4]}"
    code = f"functest_{suffix}"
    resp = client.post(
        f"{API}/b01/spaces",
        headers=admin_headers,
        json={
            "space_code": code,
            "space_name": f"功能测试空间_{suffix}",
            "customer_name": "功能测试客户",
            "source_platform": "华为云",
            "target_platform": "腾讯云",
            "owner": "功能测试负责人",
            "overall_progress": 10,
        },
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["code"] == 0, body
    space = body["data"]
    yield space

    # 清理：角色绑定 → 空间记录 → 空间库（仅允许删除 functest 前缀库，防误删）
    from app.database import SessionLocal, _server_engine, _space_engines
    from app.models.role import CommonUserRole
    from app.models.space import MigSpace

    db = SessionLocal()
    try:
        db.query(CommonUserRole).filter(CommonUserRole.space_id == space["id"]).delete()
        db.query(MigSpace).filter(MigSpace.id == space["id"]).delete()
        db.commit()
    finally:
        db.close()
    db_name = space["db_name"]
    if db_name and db_name.startswith("mig_functest"):
        engine = _space_engines.pop(db_name, None)
        if engine is not None:
            engine.dispose()
        server = _server_engine()
        try:
            with server.connect() as conn:
                conn.execute(text(f"DROP DATABASE IF EXISTS `{db_name}`"))
        finally:
            server.dispose()


@pytest.fixture(scope="session")
def member_headers(client, admin_headers, test_space):
    """空间内 L2 项目成员登录态。"""
    username = uid("functest_member")
    user_id = _create_user(client, admin_headers, username)
    resp = client.put(
        f"{API}/common/users/{user_id}/roles",
        headers=admin_headers,
        json={"roles": [{"role_level": "L2", "space_id": test_space["id"]}]},
    )
    assert_ok(resp)
    try:
        body = _login(client, username, MEMBER_PASSWORD)
        assert body["code"] == 0, body
        yield {"Authorization": f"Bearer {body['data']['token']}"}
    finally:
        _hard_delete_user(user_id)


@pytest.fixture(scope="session")
def l1_headers(client, admin_headers, test_space):
    """空间内 L1 项目管理员登录态（可创建空间）。"""
    username = uid("functest_l1")
    user_id = _create_user(client, admin_headers, username)
    resp = client.put(
        f"{API}/common/users/{user_id}/roles",
        headers=admin_headers,
        json={"roles": [{"role_level": "L1", "space_id": test_space["id"]}]},
    )
    assert_ok(resp)
    try:
        body = _login(client, username, MEMBER_PASSWORD)
        assert body["code"] == 0, body
        yield {"Authorization": f"Bearer {body['data']['token']}"}
    finally:
        _hard_delete_user(user_id)


@pytest.fixture(scope="session")
def outsider_headers(client, admin_headers):
    """无任何空间角色的平台用户登录态（越权访问验证用）。"""
    username = uid("functest_outsider")
    user_id = _create_user(client, admin_headers, username)
    try:
        body = _login(client, username, MEMBER_PASSWORD)
        assert body["code"] == 0, body
        yield {"Authorization": f"Bearer {body['data']['token']}"}
    finally:
        _hard_delete_user(user_id)


@pytest.fixture()
def captcha_helper(client):
    """获取验证码及其明文（同进程内存存储，仅供测试使用）。"""
    from app.core import security

    def _get():
        resp = client.get(f"{API}/common/auth/captcha")
        assert_ok(resp)
        captcha_id = resp.json()["data"]["captcha_id"]
        return captcha_id, security._captcha_store[captcha_id]["code"]

    return _get


@pytest.fixture()
def throwaway_user(client, admin_headers):
    """一次性测试用户（登录锁定等破坏性场景用），结束即删除。"""
    state = {}

    def _create(prefix: str = "functest_tmp"):
        username = uid(prefix)
        state["user_id"] = _create_user(client, admin_headers, username)
        state["username"] = username
        return state

    yield _create
    if state.get("user_id"):
        _hard_delete_user(state["user_id"])
